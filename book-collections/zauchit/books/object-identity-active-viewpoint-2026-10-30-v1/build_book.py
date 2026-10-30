#!/usr/bin/env python3
"""Build the illustrated Russian mini-book as PDF and DOCX from its Markdown source."""
from pathlib import Path
import re
import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from weasyprint import HTML
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent
MD_PATH = BASE / "ZA_ZASLONKOY-v1.md"
PDF_PATH = BASE / "ZA_ZASLONKOY-v1.pdf"
DOCX_PATH = BASE / "ZA_ZASLONKOY-v1.docx"

NAVY = "183241"
TEAL = "3C8178"
GOLD = "B99755"
INK = "26343D"
MUTED = "596772"


def add_hyperlink(paragraph, text, url):
    """Add an external clickable hyperlink to a Word paragraph."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GOLD)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_inline(paragraph, node, bold=False, italic=False, code=False):
    if isinstance(node, NavigableString):
        if str(node):
            run = paragraph.add_run(str(node))
            run.bold = bold
            run.italic = italic
            if code:
                run.font.name = "DejaVu Sans Mono"
                run.font.size = Pt(9)
        return
    if not isinstance(node, Tag):
        return
    name = node.name.lower()
    if name == "br":
        paragraph.add_run().add_break()
        return
    if name == "a":
        add_hyperlink(paragraph, node.get_text("", strip=False), node.get("href", ""))
        return
    next_bold = bold or name in {"strong", "b"}
    next_italic = italic or name in {"em", "i"}
    next_code = code or name == "code"
    for child in node.children:
        add_inline(paragraph, child, next_bold, next_italic, next_code)

def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    displayed = OxmlElement("w:t")
    displayed.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, displayed, end):
        run._r.append(element)
    run.font.color.rgb = RGBColor.from_string(GOLD)
    run.font.size = Pt(9)


def add_horizontal_rule(document):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), GOLD)
    borders.append(bottom)
    pPr.append(borders)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def markdown_body_and_html():
    text = MD_PATH.read_text(encoding="utf-8")
    marker = "### О чём эта книга"
    body = text[text.index(marker):]
    html_body = markdown.markdown(
        body,
        extensions=["extra", "toc", "sane_lists", "smarty"],
        output_format="html5",
    )
    soup = BeautifulSoup(html_body, "html.parser")
    for image in soup.find_all("img"):
        src = image.get("src", "")
        image["class"] = ["diagram" if src.startswith("diagrams/") else "art"]
        image["loading"] = "eager"
        parent = image.parent
        if parent and parent.name == "p" and len(parent.contents) == 1:
            figure = soup.new_tag("figure")
            figure["class"] = image["class"]
            image.extract()
            figure.append(image)
            parent.replace_with(figure)
    for paragraph in soup.find_all("p"):
        if re.match(r"^\[\d+\]", paragraph.get_text("", strip=True)):
            paragraph["class"] = list(paragraph.get("class", [])) + ["reference"]
    return soup


def build_pdf(soup):
    body_html = str(soup)
    body_html = re.sub(r'src="assets/([A-Za-z0-9_-]+)\.png"', r'src="assets/print/\1.jpg"', body_html)
    html = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><title>За заслонкой — иллюстрированная мини-книга</title>
<style>
@page {{ size: A4; margin: 18mm 19mm 22mm 19mm;
  @top-left {{ content: "ЗА ЗАСЛОНКОЙ  ·  ИДЕНТИЧНОСТЬ ОБЪЕКТА"; color: #8e784d; font: 7.5pt 'DejaVu Sans', sans-serif; letter-spacing: .4pt; }}
  @bottom-center {{ content: counter(page); color: #8e784d; font: 9pt 'DejaVu Sans', sans-serif; }}
}}
@page:first {{ margin: 0; @top-left {{ content: none; }} @bottom-center {{ content: none; }} }}
html {{ font-family: 'DejaVu Serif', Georgia, serif; color: #{INK}; font-size: 10.4pt; line-height: 1.43; }}
body {{ margin: 0; }}
.cover {{ position: relative; width: 210mm; height: 297mm; page-break-after: always; overflow: hidden; background: #101b2a; }}
.cover img {{ position: absolute; inset: 0; width: 210mm; height: 297mm; object-fit: cover; }}
.cover-copy {{ position: absolute; top: 28mm; left: 18mm; right: 18mm; text-align: center; color: #f2d587; text-shadow: 0 1px 5px #101521; }}
.cover-kicker {{ margin: 0 0 7mm; font: 9pt 'DejaVu Sans', sans-serif; letter-spacing: 2.1pt; text-transform: uppercase; }}
.cover h1 {{ page-break-before: auto; break-before: auto; margin: 0 0 5mm; font: 700 32pt 'DejaVu Serif', Georgia, serif; letter-spacing: 1.1pt; color: #f2d587; }}
.cover-subtitle {{ margin: 0 auto; max-width: 150mm; font: 15pt 'DejaVu Serif', Georgia, serif; line-height: 1.35; color: #f5e8c6; }}
.cover-edition {{ position: absolute; bottom: 25mm; left: 18mm; right: 18mm; text-align: center; color: #edd28f; font: 9pt 'DejaVu Sans', sans-serif; letter-spacing: .6pt; line-height: 1.6; text-shadow: 0 1px 4px #101521; }}
main {{ display: block; }}
h1 {{ page-break-before: always; break-before: page; color: #{NAVY}; font: 700 22pt 'DejaVu Serif', Georgia, serif; line-height: 1.2; margin: 0 0 8mm; padding: 0 0 3mm; border-bottom: 1.4pt solid #{GOLD}; }}
h2 {{ color: #{NAVY}; font: 700 14.5pt 'DejaVu Serif', Georgia, serif; line-height: 1.22; margin: 7mm 0 3mm; break-after: avoid; }}
h3 {{ color: #{TEAL}; font: 700 11.6pt 'DejaVu Sans', sans-serif; margin: 5mm 0 2mm; break-after: avoid; }}
p {{ margin: 0 0 3.4mm; orphans: 3; widows: 3; }}
a {{ color: #8c6f37; text-decoration: none; }}
strong, b {{ color: #{NAVY}; }}
blockquote {{ margin: 4mm 0; padding: 3mm 5mm; border-left: 3pt solid #{GOLD}; background: #f4f1e8; color: #{NAVY}; }}
blockquote p:last-child {{ margin-bottom: 0; }}
ul, ol {{ margin: 2mm 0 4mm; padding-left: 7mm; }}
li {{ margin: 0 0 1.8mm; padding-left: 1mm; }}
hr {{ border: 0; border-top: 1pt solid #{GOLD}; margin: 5mm 0; }}
figure {{ margin: 5mm 0 3mm; padding: 0; text-align: center; break-inside: avoid; page-break-inside: avoid; }}
figure.art img {{ display: block; margin: 0 auto; max-width: 100%; max-height: 100mm; width: auto; object-fit: contain; }}
figure.diagram img {{ display: block; margin: 0 auto; max-width: 100%; max-height: 155mm; width: auto; object-fit: contain; }}
.reference {{ font-size: 8.5pt; line-height: 1.28; margin: 0 0 2.2mm; padding-left: 9mm; text-indent: -9mm; overflow-wrap: anywhere; }}
@media print {{ h1, h2, h3, figure, blockquote {{ break-inside: avoid; }} }}
</style></head><body>
<section class="cover">
  <img src="assets/print/cover.jpg" alt="Авторская ночная сцена в декоративной золотой рамке">
  <div class="cover-copy">
    <p class="cover-kicker">Иллюстрированная научная мини-книга</p>
    <h1>ЗА ЗАСЛОНКОЙ</h1>
    <p class="cover-subtitle">Идентичность объекта,<br>смена ракурса и выбор следующего взгляда</p>
  </div>
  <p class="cover-edition">Редакция v1 · 30 октября 2026<br>Обзор источников и план проверок AVO-01—AVO-05</p>
</section>
<main>{body_html}</main>
</body></html>"""
    HTML(string=html, base_url=str(BASE)).write_pdf(str(PDF_PATH))


def add_image(document, image_tag):
    src = image_tag.get("src", "")
    if src.startswith("assets/"):
        optimized = BASE / "assets" / "print" / (Path(src).stem + ".jpg")
        path = optimized if optimized.exists() else (BASE / src)
    else:
        path = BASE / src
    path = path.resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(6.1))
    return p


def add_list(document, list_tag, ordered=False):
    for li in list_tag.find_all("li", recursive=False):
        style = "List Number" if ordered else "List Bullet"
        p = document.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        for child in li.children:
            if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                continue
            if isinstance(child, Tag) and child.name == "p":
                for sub in child.children:
                    add_inline(p, sub)
            else:
                add_inline(p, child)
        for nested in li.find_all(["ul", "ol"], recursive=False):
            add_list(document, nested, ordered=nested.name == "ol")


def build_docx(soup):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(19)
    section.right_margin = Mm(19)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Serif"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, font in [
        ("Heading 1", 22, NAVY, "DejaVu Serif"),
        ("Heading 2", 14, NAVY, "DejaVu Serif"),
        ("Heading 3", 11, TEAL, "DejaVu Sans"),
    ]:
        style = doc.styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(11 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(4)

    # A dedicated illustrated title page.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run("ЗА ЗАСЛОНКОЙ")
    r.bold = True
    r.font.name = "DejaVu Serif"
    r.font.size = Pt(29)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Идентичность объекта, смена ракурса\nи выбор следующего взгляда")
    r.font.name = "DejaVu Serif"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    cover = BASE / "assets" / "print" / "cover.jpg"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(cover), width=Inches(4.8))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Иллюстрированная научная мини-книга · редакция v1\n30 октября 2026")
    r.font.name = "DejaVu Sans"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    doc.add_page_break()

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("ЗА ЗАСЛОНКОЙ  ·  ИДЕНТИЧНОСТЬ ОБЪЕКТА")
    header_run.font.name = "DejaVu Sans"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(GOLD)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Иллюстрированная редакция v1  ·  ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_field(footer)

    for block in list(soup.contents):
        if isinstance(block, NavigableString) or not isinstance(block, Tag):
            continue
        name = block.name.lower()
        if name in {"h1", "h2", "h3"}:
            level = int(name[1])
            p = doc.add_paragraph(style=f"Heading {level}")
            if level == 1:
                p.paragraph_format.page_break_before = True
            for child in block.children:
                add_inline(p, child)
        elif name == "p":
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            if "reference" in block.get("class", []):
                p.paragraph_format.left_indent = Inches(0.32)
                p.paragraph_format.first_line_indent = Inches(-0.32)
                p.paragraph_format.space_after = Pt(3)
                for child in block.children:
                    add_inline(p, child)
                for run in p.runs:
                    run.font.size = Pt(8.5)
            else:
                for child in block.children:
                    add_inline(p, child)
        elif name == "figure":
            image = block.find("img")
            if image:
                add_image(doc, image)
        elif name in {"ul", "ol"}:
            add_list(doc, block, ordered=name == "ol")
        elif name == "blockquote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.right_indent = Inches(0.12)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F4F1E8")
            pPr.append(shd)
            for child in block.children:
                add_inline(p, child)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(NAVY)
        elif name == "hr":
            add_horizontal_rule(doc)
        elif name == "table":
            rows = block.find_all("tr")
            if rows:
                cols = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
                table = doc.add_table(rows=0, cols=cols)
                table.style = "Light Shading Accent 1"
                for row_idx, row in enumerate(rows):
                    cells = table.add_row().cells
                for j, cell_tag in enumerate(row.find_all(["th", "td"], recursive=False)):
                    cells[j].text = cell_tag.get_text(" ", strip=True)
                    if row_idx == 0:
                        set_cell_shading(cells[j], "E9E2D1")

    doc.core_properties.title = "За заслонкой: идентичность объекта, смена ракурса и активный выбор взгляда"
    doc.core_properties.subject = "Иллюстрированный обзор исследований и пять предложенных гипотез AVO"
    doc.core_properties.language = "ru-RU"
    doc.save(DOCX_PATH)


def main():
    soup = markdown_body_and_html()
    build_pdf(soup)
    build_docx(soup)
    print(f"PDF: {PDF_PATH} ({PDF_PATH.stat().st_size:,} bytes)")
    print(f"DOCX: {DOCX_PATH} ({DOCX_PATH.stat().st_size:,} bytes)")
    print(f"Markdown: {MD_PATH} ({MD_PATH.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
